import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(summary_df, fields_details_dict, gdb_name="Geodatabase"):
    """
    Genera un informe Excel estructurado con openpyxl conteniendo:
    - Hoja 1: Resumen_GDB (Salud de capas y métricas generales)
    - Hoja 2: Campos_Incompletos (Consolidado de campos a completar por el usuario)
    - Hoja 3: Estructura_Completa (Inventario técnico de todos los campos de la GDB)
    """
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    
    # Estilos
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    
    warn_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Salmón suave
    ok_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")   # Verde suave
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # -------------------------------------------------------------
    # HOJA 1: RESUMEN_GDB
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Resumen_GDB"
    
    ws1["A1"] = f"Reporte de Revisión y Auditoría de Geodatabase: {gdb_name}"
    ws1["A1"].font = title_font
    ws1["A2"] = "Resumen ejecutivo de Feature Classes, estructura y diagnóstico de valores nulos."
    ws1["A2"].font = subtitle_font
    
    # Escribir encabezados de tabla a partir de la fila 4
    start_row = 4
    headers_1 = list(summary_df.columns)
    for col_num, h_text in enumerate(headers_1, 1):
        cell = ws1.cell(row=start_row, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for r_idx, row_data in enumerate(summary_df.itertuples(index=False), start_row + 1):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            
            col_name = headers_1[c_idx - 1]
            if col_name == '% Completitud Capa':
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '0.00"%"'
                if isinstance(val, (int, float)):
                    cell.fill = ok_fill if val >= 100.0 else warn_fill
            elif isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    # -------------------------------------------------------------
    # HOJA 2: CAMPOS_INCOMPLETOS
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Campos_Incompletos")
    ws2["A1"] = "Consolidado de Campos con Valores Vacíos (Pendientes por Completar)"
    ws2["A1"].font = title_font
    ws2["A2"] = "Lista prioritaria de Feature Classes y atributos que requieren revisión y llenado de información."
    ws2["A2"].font = subtitle_font
    
    headers_2 = ["Feature Class", "Campo", "Tipo de Dato", "Cant. Vacíos", "% Vacíos", "Cant. Llenos", "% Completitud", "Valores Únicos", "Ejemplos de Datos"]
    for col_num, h_text in enumerate(headers_2, 1):
        cell = ws2.cell(row=4, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    row_counter = 5
    for fc_name, fields_df in fields_details_dict.items():
        incomplete_fields = fields_df[fields_df['Cant. Vacíos'] > 0]
        for _, row in incomplete_fields.iterrows():
            ws2.cell(row=row_counter, column=1, value=fc_name).border = thin_border
            ws2.cell(row=row_counter, column=2, value=row['Campo']).border = thin_border
            ws2.cell(row=row_counter, column=3, value=row['Tipo de Dato']).border = thin_border
            
            c4 = ws2.cell(row=row_counter, column=4, value=row['Cant. Vacíos'])
            c4.border = thin_border
            c4.alignment = Alignment(horizontal="right")
            c4.fill = warn_fill
            
            c5 = ws2.cell(row=row_counter, column=5, value=row['% Vacíos'])
            c5.border = thin_border
            c5.alignment = Alignment(horizontal="right")
            c5.number_format = '0.00"%"'
            
            c6 = ws2.cell(row=row_counter, column=6, value=row['Cant. Llenos'])
            c6.border = thin_border
            c6.alignment = Alignment(horizontal="right")
            
            c7 = ws2.cell(row=row_counter, column=7, value=row['% Completitud'])
            c7.border = thin_border
            c7.alignment = Alignment(horizontal="right")
            c7.number_format = '0.00"%"'
            
            ws2.cell(row=row_counter, column=8, value=row['Valores Únicos']).border = thin_border
            ws2.cell(row=row_counter, column=9, value=row['Ejemplos de Datos']).border = thin_border
            row_counter += 1

    if row_counter == 5:
        ws2.cell(row=5, column=1, value="¡Felicitaciones! No se encontraron campos con valores vacíos en ninguna Feature Class.").font = Font(bold=True, color="385723")

    # -------------------------------------------------------------
    # HOJA 3: ESTRUCTURA_COMPLETA
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Estructura_Completa")
    ws3["A1"] = "Inventario y Estructura Completa de Atributos de la GDB"
    ws3["A1"].font = title_font
    ws3["A2"] = "Detalle exhaustivo de todos los campos existentes en cada Feature Class."
    ws3["A2"].font = subtitle_font
    
    headers_3 = ["Feature Class", "Campo", "Tipo de Dato", "Estado", "Cant. Vacíos", "% Vacíos", "Cant. Llenos", "% Completitud", "Valores Únicos", "Ejemplos de Datos"]
    for col_num, h_text in enumerate(headers_3, 1):
        cell = ws3.cell(row=4, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    row_counter_3 = 5
    for fc_name, fields_df in fields_details_dict.items():
        for _, row in fields_df.iterrows():
            ws3.cell(row=row_counter_3, column=1, value=fc_name).border = thin_border
            ws3.cell(row=row_counter_3, column=2, value=row['Campo']).border = thin_border
            ws3.cell(row=row_counter_3, column=3, value=row['Tipo de Dato']).border = thin_border
            
            c_estado = ws3.cell(row=row_counter_3, column=4, value=row['Estado'])
            c_estado.border = thin_border
            c_estado.alignment = Alignment(horizontal="center")
            c_estado.fill = ok_fill if row['Estado'] == 'COMPLETO' else warn_fill
            
            ws3.cell(row=row_counter_3, column=5, value=row['Cant. Vacíos']).border = thin_border
            
            c_vpct = ws3.cell(row=row_counter_3, column=6, value=row['% Vacíos'])
            c_vpct.border = thin_border
            c_vpct.number_format = '0.00"%"'
            
            ws3.cell(row=row_counter_3, column=7, value=row['Cant. Llenos']).border = thin_border
            
            c_cpct = ws3.cell(row=row_counter_3, column=8, value=row['% Completitud'])
            c_cpct.border = thin_border
            c_cpct.number_format = '0.00"%"'
            
            ws3.cell(row=row_counter_3, column=9, value=row['Valores Únicos']).border = thin_border
            ws3.cell(row=row_counter_3, column=10, value=row['Ejemplos de Datos']).border = thin_border
            row_counter_3 += 1

    # Ajuste automático del ancho de columnas en todas las hojas
    for ws in [ws1, ws2, ws3]:
        ws.row_dimensions[4].height = 28
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row < 4:
                    continue  # Ignorar títulos para no distorsionar el ancho
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output)
    output.seek(0)
    return output.getvalue()
